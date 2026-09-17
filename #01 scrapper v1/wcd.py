import asyncio
import os
import re
from pathlib import Path
from urllib.parse import urljoin

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from playwright.async_api import async_playwright


SCRIPT_DIR = Path(__file__).resolve().parent
ENV_FILE = SCRIPT_DIR / ".env"
if ENV_FILE.exists():
    for line in ENV_FILE.read_text(encoding="utf-8").splitlines():
        if "=" in line and not line.strip().startswith("#"):
            key, value = line.split("=", 1)
            os.environ.setdefault(key.strip(), value.strip().strip('"'))

BASE = Path(os.environ.get("DOWNLOAD_BASE_PATH", SCRIPT_DIR)) / "WCD"
AUDIT_FILE = BASE / "WCD_Audit.xlsx"
LEGISLATIONS_URL = "https://wcd.gov.in/documents/legislations"
ORDERS_URL = "https://wcd.gov.in/documents/orders-and-notices"
HEADLESS = os.environ.get("WCD_HEADLESS", "false").lower() in {"1", "true", "yes"}
MAX_ITEMS = int(os.environ.get("WCD_MAX_ITEMS", "0"))  # 0 = full run


def safe_filename(value):
    value = re.sub(r"\s+", " ", value).strip()
    value = re.sub(r'[<>:"/\\|?*]', "_", value)
    return value[:170] or "WCD_document"


def unique_path(folder, filename):
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / filename
    n = 1
    while path.exists():
        path = folder / f"{Path(filename).stem}_{n}{Path(filename).suffix}"
        n += 1
    return path


def write_audit(rows):
    BASE.mkdir(parents=True, exist_ok=True)
    if AUDIT_FILE.exists():
        workbook = load_workbook(AUDIT_FILE)
        if "Audit" in workbook.sheetnames:
            del workbook["Audit"]
    else:
        workbook = Workbook()
        if "Sheet" in workbook.sheetnames:
            del workbook["Sheet"]
    sheet = workbook.create_sheet("Audit")
    sheet.append(["Name", "Category", "Published Date", "PDF Path"])
    for row in rows:
        sheet.append([row["name"], row["category"], row["date"], row["pdf_path"]])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F5597")
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for col, width in {"A": 110, "B": 24, "C": 20, "D": 115}.items():
        sheet.column_dimensions[col].width = width
    workbook.save(AUDIT_FILE)


def posh_match(title):
    text = title.lower()
    return "sexual harassment" in text and ("workplace" in text or "women" in text)


async def download_pdf(context, url, folder, number, title):
    try:
        response = await context.request.get(url, timeout=60000)
        data = await response.body()
        if not response.ok or not data.startswith(b"%PDF"):
            return f"HTTP {response.status} / NOT PDF"
        path = unique_path(folder, f"{number:03d}_{safe_filename(title)}.pdf")
        path.write_bytes(data)
        return str(path.resolve())
    except Exception as exc:
        return f"ERROR: {type(exc).__name__}: {exc}"


async def extract_page_items(page, section):
    return await page.evaluate("""section => {
      const out = [];
      const links = [...document.querySelectorAll('a[href]')];
      for (const link of links) {
        const title = (link.getAttribute('aria-label') || link.innerText || '').replace(/\\s+/g, ' ').trim();
        if (!title || title.toLowerCase() === 'view' || !/documents\\/uploaded\\/.*\\.pdf/i.test(link.href)) continue;
        let block = link.parentElement, text = '';
        for (let i = 0; block && i < 5; i++, block = block.parentElement) {
          text = (block.innerText || '').replace(/\\s+/g, ' ').trim();
          if (text.length > title.length) break;
        }
        const date = (text.match(/\\b\\d{2}[-/]\\d{2}[-/]\\d{4}\\b/) || [''])[0];
        const type = section === 'Orders and Notices' ? section : (text.match(/\\b(ACTs|Rules)\\b/i) || ['Document'])[0];
        out.push({name: title.replace(/\\s*-\\s*PDF file.*$/i, '').trim(), category: type, date, url: link.href});
      }
      return out;
    }""", section)


async def collect(context, page, base_url, section, folder, rows):
    page_number = 1
    seen = set()
    while True:
        await page.goto(f"{base_url}?page={page_number}" if page_number > 1 else base_url,
                        wait_until="domcontentloaded", timeout=60000)
        await page.wait_for_timeout(1800)
        items = await extract_page_items(page, section)
        for item in items:
            if not posh_match(item["name"]) or item["url"] in seen:
                continue
            seen.add(item["url"])
            item["pdf_path"] = await download_pdf(context, item["url"], folder, len(seen), item["name"])
            rows.append(item)
            write_audit(rows)
            print(f"{section} {len(seen)}: {item['name']} -> {item['pdf_path']}", flush=True)
            if MAX_ITEMS and len(rows) >= MAX_ITEMS:
                return
        next_url = await page.locator("a", has_text=re.compile(r"Next", re.I)).last.get_attribute("href") if await page.locator("a", has_text=re.compile(r"Next", re.I)).count() else None
        if not next_url:
            return
        page_number += 1


async def main():
    rows = []
    write_audit(rows)
    async with async_playwright() as p:
        browser = await p.chromium.launch(headless=HEADLESS)
        context = await browser.new_context()
        page = await context.new_page()
        await collect(context, page, LEGISLATIONS_URL, "Acts and Rules", BASE / "Acts_and_Rules", rows)
        if not MAX_ITEMS or len(rows) < MAX_ITEMS:
            await collect(context, page, ORDERS_URL, "Orders and Notices", BASE / "Orders_and_Notices", rows)
        await context.close()
        await browser.close()
    print(f"POSH records: {len(rows)}")
    print(f"Audit: {AUDIT_FILE.resolve()}")


if __name__ == "__main__":
    asyncio.run(main())
