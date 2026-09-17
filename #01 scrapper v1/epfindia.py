from datetime import datetime
from pathlib import Path
from urllib.parse import urljoin
import os
import re
import time

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from playwright.sync_api import sync_playwright


SCRIPT_DIR = Path(__file__).resolve().parent


def load_env_file():
    env_path = SCRIPT_DIR / ".env"
    if not env_path.exists():
        return
    for raw in env_path.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        os.environ.setdefault(key.strip(), value.strip().strip('"'))


load_env_file()

DOWNLOAD_BASE = Path(os.environ.get("DOWNLOAD_BASE_PATH", SCRIPT_DIR))
BASE = DOWNLOAD_BASE / "EPFO"
ACTS_FOLDER = BASE / "Acts"
CIRCULARS_FOLDER = BASE / "Circulars"
AUDIT_XLSX = BASE / "EPFO_Audit.xlsx"

ACT_URL = "https://www.epfo.gov.in/epf-mp-act-1952/"
CIRCULARS_URL = "https://www.epfo.gov.in/circulars/"
TIMEOUT = 60000
DOWNLOAD_PDFS = os.environ.get("EPFO_DOWNLOAD_PDFS", "true").lower() not in {"0", "false", "no"}
HEADLESS = os.environ.get("EPFO_HEADLESS", "false").lower() not in {"0", "false", "no"}


def safe_filename(value):
    value = re.sub(r'[<>:"/\\|?*]', "_", value)
    value = re.sub(r"\s+", " ", value).strip()
    return (value[:170] or "EPFO_document")


def unique_path(folder, filename):
    folder.mkdir(parents=True, exist_ok=True)
    candidate = folder / filename
    n = 1
    while candidate.exists():
        candidate = folder / f"{Path(filename).stem}_{n}{Path(filename).suffix}"
        n += 1
    return candidate


def extract_circulars(page):
    page.goto(CIRCULARS_URL, wait_until="domcontentloaded", timeout=TIMEOUT)
    page.wait_for_timeout(5000)
    return page.evaluate("""() => {
      const dateRe = /\\b\\d{2}\\/\\d{2}\\/\\d{4}\\b/;
      const bad = new Set(['View PDF', 'Circulars']);
      const rows = [];
      for (const link of [...document.querySelectorAll('a')].filter(a => a.innerText.trim().toLowerCase() === 'view pdf')) {
        let node = link, best = '';
        for (let i = 0; node && i < 9; i++, node = node.parentElement) {
          const text = (node.innerText || '').trim();
          if (dateRe.test(text) && text.length > best.length && text.length < 2500) best = text;
        }
        const lines = best.split(/\\n+/).map(x => x.trim()).filter(Boolean);
        const date = (best.match(dateRe) || [''])[0];
        const dateIndex = lines.findIndex(x => dateRe.test(x));
        const beforeDate = lines.slice(0, dateIndex < 0 ? lines.length : dateIndex);
        const name = beforeDate.filter(x => !/^\\d+\\.$/.test(x) && !bad.has(x)).sort((a,b) => b.length-a.length)[0] || 'Unnamed circular';
        const category = (beforeDate.find(x => /^[A-Za-z&\\s]+$/.test(x) && x.length < 60 && x !== name) || 'Circular')
          .replace(/^[•·]\\s*/, '').trim();
        const href = link.href && !link.href.includes('/false') ? link.href : '';
        rows.push({name, category, date, href});
      }
      return rows;
    }""")


def download_pdf(context, url, folder, index, name):
    if not url:
        return "NOT AVAILABLE"
    try:
        response = context.request.get(url, timeout=TIMEOUT)
        content = response.body()
        if not response.ok or not content.startswith(b"%PDF"):
            return "DOWNLOAD FAILED"
        filename = f"{index:04d}_{safe_filename(name)}.pdf"
        path = unique_path(folder, filename)
        path.write_bytes(content)
        return str(path.resolve())
    except Exception:
        return "DOWNLOAD FAILED"


def write_audit(rows):
    BASE.mkdir(parents=True, exist_ok=True)
    if AUDIT_XLSX.exists():
        workbook = load_workbook(AUDIT_XLSX)
        if "Audit" in workbook.sheetnames:
            del workbook["Audit"]
    else:
        workbook = Workbook()
        if "Sheet" in workbook.sheetnames:
            del workbook["Sheet"]
    sheet = workbook.create_sheet("Audit")
    headers = ["Name", "Category", "Date", "PDF Path"]
    sheet.append(headers)
    for row in rows:
        sheet.append([row["name"], row["category"], row["date"], row["pdf_path"]])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F5597")
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = {"A": 75, "B": 24, "C": 14, "D": 110}
    for col, width in widths.items():
        sheet.column_dimensions[col].width = width
    workbook.save(AUDIT_XLSX)


def main():
    ACTS_FOLDER.mkdir(parents=True, exist_ok=True)
    CIRCULARS_FOLDER.mkdir(parents=True, exist_ok=True)
    rows = [{
        "name": "Employees' Provident Funds and Miscellaneous Provisions Act, 1952",
        "category": "Act",
        "date": "",
        "pdf_path": str((ACTS_FOLDER / "EPF_MP_Act_1952.pdf").resolve()),
    }]
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=HEADLESS)
        page = browser.new_page(viewport={"width": 1366, "height": 900})
        page.goto(ACT_URL, wait_until="networkidle", timeout=TIMEOUT)
        page.pdf(path=rows[0]["pdf_path"], format="A4", print_background=True)
        circulars = extract_circulars(page)
        for index, item in enumerate(circulars, 1):
            existing = sorted(CIRCULARS_FOLDER.glob(f"{index:04d}_*.pdf"))
            if existing:
                item["pdf_path"] = str(existing[0].resolve())
            else:
                item["pdf_path"] = download_pdf(page.context, item["href"], CIRCULARS_FOLDER, index, item["name"]) if DOWNLOAD_PDFS else "NOT DOWNLOADED"
            rows.append(item)
        browser.close()
    write_audit(rows)
    valid = sum(1 for row in rows[1:] if row["pdf_path"].endswith(".pdf"))
    unavailable = sum(1 for row in rows[1:] if row["pdf_path"] == "NOT AVAILABLE")
    failed = sum(1 for row in rows[1:] if row["pdf_path"] == "DOWNLOAD FAILED")
    print(f"Act records       : {len(rows) - len(circulars)}")
    print(f"Circular records  : {len(circulars)}")
    print(f"Valid PDFs        : {valid}")
    print(f"Unavailable PDFs  : {unavailable}")
    print(f"Failed downloads  : {failed}")
    print(f"Audit workbook    : {AUDIT_XLSX.resolve()}")


if __name__ == "__main__":
    main()
