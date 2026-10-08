
"""
Provident Fund Act 1952 / EPFO document downloader
===================================================

Run from PowerShell:
    python "006 Provident Fund Act 1952.py"

Execution order and output:
    1. EPF_MP_Act_1952_PDFs/  - EPF Act webpage PDF
    2. EPFO_Circulars_PDFs/    - all EPFO circular PDFs
    3. EPFO_Circular_Report.xlsx - download audit report
"""

from pathlib import Path
import os
import json
import random
from urllib.parse import urljoin, urlparse
import re
import time
from datetime import datetime
from collections import Counter

from playwright.sync_api import sync_playwright
from openpyxl import Workbook, load_workbook


# ============================================================
# CONFIGURATION
# ============================================================

BASE_URL = "https://www.epfo.gov.in"

START_URL = "https://www.epfo.gov.in/"

EPF_ACT_URL = "https://www.epfo.gov.in/epf-mp-act-1952/"

CIRCULARS_URL = "https://www.epfo.gov.in/circulars/"

SCRIPT_DIR = Path(__file__).resolve().parent
ENV_FILE = next(parent / ".env" for parent in Path(__file__).resolve().parents if (parent / "docker-compose.yml").is_file())
if ENV_FILE.exists():
    for line in ENV_FILE.read_text(encoding="utf-8").splitlines():
        if "=" in line and not line.lstrip().startswith("#"):
            key, value = line.split("=", 1)
            os.environ.setdefault(key.strip(), value.strip().strip('"'))
DOWNLOAD_ROOT = Path(os.environ.get("DOWNLOAD_BASE_PATH", SCRIPT_DIR)) / Path(__file__).stem
RUNTIME_DIR = DOWNLOAD_ROOT / "logs"
RUN_STATUS = RUNTIME_DIR / "run_status.json"
RUNTIME_LOG = RUNTIME_DIR / "runtime.log"
RESILIENCE_ATTEMPTS = int(os.environ.get("SCRAPER_MAX_ATTEMPTS", "3"))
REQUEST_DELAY_SECONDS = float(os.environ.get("SCRAPER_REQUEST_DELAY_SECONDS", "1.0"))
BLOCK_COOLDOWN = float(os.environ.get("SCRAPER_BLOCK_COOLDOWN_SECONDS", "300"))


class SiteAccessBlocked(RuntimeError):
    pass


BLOCK_DETECTED = ""


def check_response_status(status, url=""):
    global BLOCK_DETECTED
    if status in {403, 429}:
        BLOCK_DETECTED = f"HTTP {status} from {url}"
        raise SiteAccessBlocked(BLOCK_DETECTED)


def write_run_status(state, message="", attempt=0):
    RUNTIME_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().astimezone().isoformat(timespec="seconds")
    payload = {"script": Path(__file__).name, "state": state, "attempt": attempt,
               "message": str(message), "timestamp": stamp}
    temporary = RUN_STATUS.with_suffix(".tmp")
    temporary.write_text(json.dumps(payload, indent=2), encoding="utf-8")
    temporary.replace(RUN_STATUS)
    with RUNTIME_LOG.open("a", encoding="utf-8") as stream:
        stream.write(f"{stamp} | {state} | attempt={attempt} | {message}\n")


def guarded_main(entrypoint):
    global BLOCK_DETECTED
    for attempt in range(1, RESILIENCE_ATTEMPTS + 1):
        BLOCK_DETECTED = ""
        write_run_status("running", attempt=attempt)
        try:
            result = entrypoint()
            if BLOCK_DETECTED:
                raise SiteAccessBlocked(BLOCK_DETECTED)
            write_run_status("completed", attempt=attempt)
            return result
        except KeyboardInterrupt:
            write_run_status("stopped", "Stopped by user", attempt)
            raise
        except SiteAccessBlocked as error:
            write_run_status("blocked", error, attempt)
            if attempt == RESILIENCE_ATTEMPTS:
                raise
            time.sleep(BLOCK_COOLDOWN + random.uniform(0, 5))
        except Exception as error:
            write_run_status("error", error, attempt)
            if attempt == RESILIENCE_ATTEMPTS:
                raise
            time.sleep(min(120, 5 * (2 ** (attempt - 1))) + random.uniform(0, 2))
EPF_PDF_FOLDER = DOWNLOAD_ROOT / "EPF_MP_Act_1952_PDFs"
CIRCULARS_FOLDER = DOWNLOAD_ROOT / "EPFO_Circulars_PDFs"
EXCEL_FILE = DOWNLOAD_ROOT / "EPFO_Circular_Report.xlsx"

HEADLESS = False
TIMEOUT = 60000


# ============================================================
# SAFE FILENAME
# ============================================================

def safe_filename(name):

    name = str(name or "").strip()

    name = re.sub(r'[<>:"/\\|?*]', "_", name)
    name = re.sub(r"\s+", " ", name)
    name = name.strip(" .")

    if not name:
        name = "EPFO_Circular"

    return name[:180]


# ============================================================
# URL FALLBACK NAME
# ============================================================

def get_filename_from_url(pdf_url):

    filename = Path(
        urlparse(pdf_url).path
    ).name

    if filename.lower().endswith(".pdf"):
        filename = filename[:-4]

    filename = filename.replace("_", " ")
    filename = filename.replace("-", " ")

    return safe_filename(filename)


# ============================================================
# UNIQUE FILE PATH
# ============================================================

def unique_file_path(folder, title, used_names):

    folder.mkdir(
        parents=True,
        exist_ok=True
    )

    title = safe_filename(title)

    count = used_names[title]
    used_names[title] += 1

    if count == 0:
        filename = f"{title}.pdf"
    else:
        filename = f"{title} {count}.pdf"

    path = folder / filename

    number = count

    while path.exists():

        number += 1
        path = folder / f"{title} {number}.pdf"

    return path


# ============================================================
# OPEN LEGAL FRAMEWORK
# ============================================================

def open_legal_framework(page):

    print("\n[STEP] Opening Legal Framework...")

    selectors = [
        "a#menu-item-dropdown-8400",
        'a[title="Legal Framework"]',
        'a:has-text("Legal Framework")'
    ]

    for selector in selectors:

        try:

            locator = page.locator(selector).first

            if locator.count() > 0:

                locator.hover()
                time.sleep(1)

                print("[OK] Legal Framework opened")
                return True

        except Exception:
            continue

    print("[ERROR] Legal Framework not found")
    return False


# ============================================================
# OPEN EPF & MP ACT 1952
# ============================================================

def open_epf_act(page):

    print("\n[STEP] Opening EPF & MP Act 1952...")

    selectors = [
        'a[href="/epf-mp-act-1952"]',
        'a[title="EPF & MP Act 1952"]',
        'a:has-text("EPF & MP Act 1952")'
    ]

    for selector in selectors:

        try:

            locator = page.locator(selector).first

            if locator.count() > 0:

                locator.click()

                page.wait_for_load_state(
                    "domcontentloaded",
                    timeout=TIMEOUT
                )

                time.sleep(2)

                print("[OK] EPF & MP Act 1952 opened")
                return True

        except Exception:
            continue

    try:

        print("[INFO] Opening EPF Act directly...")

        page.goto(
            EPF_ACT_URL,
            wait_until="domcontentloaded",
            timeout=TIMEOUT
        )

        time.sleep(2)

        print("[OK] EPF & MP Act 1952 opened")
        return True

    except Exception as error:

        print(
            "[ERROR] EPF Act could not be opened:",
            error
        )

        return False


# ============================================================
# SAVE EPF ACT WEBPAGE AS PDF - ONLY ONCE
# ============================================================

def save_epf_act_as_pdf():

    EPF_PDF_FOLDER.mkdir(
        parents=True,
        exist_ok=True
    )

    pdf_path = (
        EPF_PDF_FOLDER /
        "EPF_MP_Act_1952.pdf"
    )

    if pdf_path.exists() and pdf_path.stat().st_size > 0:

        print("\n[SKIP] EPF Act PDF already exists")
        print(pdf_path.resolve())

        return True

    print("\n============================================================")
    print(" SAVING EPF & MP ACT 1952 AS PDF")
    print("============================================================")

    try:

        with sync_playwright() as p:

            browser = p.chromium.launch(
                headless=True
            )

            page = browser.new_page(
                viewport={
                    "width": 1366,
                    "height": 900
                }
            )

            page.goto(
                EPF_ACT_URL,
                wait_until="networkidle",
                timeout=TIMEOUT
            )

            time.sleep(3)

            page.pdf(
                path=str(pdf_path.resolve()),
                format="A4",
                print_background=True,
                display_header_footer=False,
                margin={
                    "top": "15mm",
                    "bottom": "15mm",
                    "left": "12mm",
                    "right": "12mm"
                }
            )

            browser.close()

        if pdf_path.exists() and pdf_path.stat().st_size > 0:

            print("[SUCCESS] EPF Act PDF saved")
            print(pdf_path.resolve())

            return True

        print("[ERROR] EPF Act PDF was not created")
        return False

    except Exception as error:

        print("[ERROR] EPF Act PDF error:", error)
        return False


# ============================================================
# OPEN CIRCULARS
# ============================================================

def open_circulars(page):

    print("\n[STEP] Opening Circulars...")

    try:

        page.goto(
            CIRCULARS_URL,
            wait_until="networkidle",
            timeout=TIMEOUT
        )

        time.sleep(5)

        # Scroll to load dynamic content
        for _ in range(15):
            page.mouse.wheel(0, 1500)
            time.sleep(0.5)

        time.sleep(3)

        print("[OK] Circulars opened")
        return True

    except Exception as error:

        print(
            "[WARNING] Direct Circulars opening failed:",
            error
        )

    return False


# ============================================================
# GET TITLE FROM PDF LINK
# ============================================================

def get_title_from_pdf_link(anchor):

    # Check attributes first
    for attribute in [
        "title",
        "aria-label",
        "data-title"
    ]:

        try:

            value = anchor.get_attribute(attribute)

            if value:

                value = value.strip()

                if (
                    len(value) > 8
                    and "download" not in value.lower()
                    and "view pdf" not in value.lower()
                    and "click here" not in value.lower()
                ):

                    return safe_filename(value)

        except Exception:
            pass

    # Search nearby parent containers
    parent = anchor

    for _ in range(10):

        try:
            parent = parent.locator("..")
        except Exception:
            break

        selectors = [
            ".views-field-title .field-content",
            ".views-field-title",
            ".field--name-title",
            ".card-title",
            ".circular-title",
            ".notice-title",
            ".document-title",
            "h1",
            "h2",
            "h3",
            "h4",
            "h5",
            "h6",
            "strong",
            "b"
        ]

        for selector in selectors:

            try:

                elements = parent.locator(selector)
                count = elements.count()

                for i in range(count):

                    text = elements.nth(i).inner_text()
                    text = re.sub(
                        r"\s+",
                        " ",
                        text
                    ).strip()

                    if len(text) < 10:
                        continue

                    if len(text) > 250:
                        continue

                    lower_text = text.lower()

                    ignored = [
                        "download",
                        "download pdf",
                        "view pdf",
                        "click here",
                        "read more",
                        "view more",
                        "circulars"
                    ]

                    if lower_text in ignored:
                        continue

                    if text.lower().endswith(".pdf"):
                        continue

                    if re.fullmatch(
                        r"[\d\s./:-]+",
                        text
                    ):
                        continue

                    return safe_filename(text)

            except Exception:
                continue

        # Get text from parent card
        try:

            card_text = parent.inner_text()

            lines = [
                re.sub(
                    r"\s+",
                    " ",
                    line
                ).strip()
                for line in card_text.splitlines()
            ]

            for line in lines:

                if len(line) < 10:
                    continue

                if len(line) > 250:
                    continue

                lower_line = line.lower()

                if any(
                    word in lower_line
                    for word in [
                        "download",
                        "view pdf",
                        "click here",
                        "read more",
                        "view more"
                    ]
                ):
                    continue

                if line.startswith("http"):
                    continue

                if re.fullmatch(
                    r"[\d\s./:-]+",
                    line
                ):
                    continue

                return safe_filename(line)

        except Exception:
            pass

    return ""


# ============================================================
# GET ALL PDF LINKS FROM CURRENT PAGE
# ============================================================

def get_pdf_links(page):

    print("\n[STEP] Searching for PDF links...")

    pdf_items = []

    # Scroll page completely
    for _ in range(15):
        page.mouse.wheel(0, 1500)
        time.sleep(0.5)

    time.sleep(3)

    anchors = page.locator("a[href]")

    total = anchors.count()

    print(f"[INFO] Total links found: {total}")

    for i in range(total):

        try:

            anchor = anchors.nth(i)

            href = anchor.get_attribute("href")

            if not href:
                continue

            full_url = urljoin(
                page.url,
                href
            )

            lower_url = full_url.lower()

            if (
                ".pdf" not in lower_url
                and "download" not in lower_url
                and "document" not in lower_url
                and "pdf" not in anchor.inner_text().lower()
            ):
                continue

            title = get_title_from_pdf_link(anchor)

            if not title:
                title = get_filename_from_url(full_url)

            item = {
                "title": safe_filename(title),
                "url": full_url
            }

            # Avoid duplicate URLs
            if not any(
                old["url"] == full_url
                for old in pdf_items
            ):

                pdf_items.append(item)

                print(
                    f"[PDF {len(pdf_items)}] "
                    f"{item['title'][:100]}"
                )

        except Exception:
            continue

    print(
        f"\n[INFO] PDFs found on current page: "
        f"{len(pdf_items)}"
    )

    return pdf_items


# ============================================================
# FIND NEXT BUTTON
# ============================================================

def find_next_button(page):

    selectors = [
        'a[rel="next"]',
        ".pager__item--next a",
        ".pagination .next a",
        "li.next a",
        "a.next",
        "a[title*='Next']",
        "a[aria-label*='Next']",
        "a:has-text('Next')",
        "button:has-text('Next')",
        "a:has-text('›')",
        "a:has-text('»')"
    ]

    for selector in selectors:

        try:

            buttons = page.locator(selector)
            count = buttons.count()

            for i in range(count):

                button = buttons.nth(i)

                if not button.is_visible():
                    continue

                class_name = (
                    button.get_attribute("class")
                    or ""
                )

                aria_disabled = button.get_attribute(
                    "aria-disabled"
                )

                if aria_disabled == "true":
                    continue

                if "disabled" in class_name.lower():
                    continue

                return button

        except Exception:
            continue

    return None


# ============================================================
# COLLECT ALL CIRCULAR PDF LINKS
# ============================================================

def collect_all_circulars(page):

    all_items = []
    visited_pages = set()
    page_number = 1

    while True:

        current_url = page.url

        if current_url in visited_pages:
            print("[INFO] Page already visited")
            break

        visited_pages.add(current_url)

        print("\n============================================================")
        print(f" READING CIRCULAR PAGE {page_number}")
        print("============================================================")

        current_items = get_pdf_links(page)

        all_items.extend(current_items)

        next_button = find_next_button(page)

        if not next_button:

            print("[INFO] Next button not found")
            break

        old_url = page.url

        try:

            next_button.scroll_into_view_if_needed()
            next_button.click()

            time.sleep(4)

            try:

                page.wait_for_load_state(
                    "domcontentloaded",
                    timeout=20000
                )

            except Exception:
                pass

            time.sleep(3)

            new_url = page.url

            if new_url == old_url:

                print(
                    "[INFO] Next page URL did not change"
                )

                break

            page_number += 1

        except Exception as error:

            print(
                "[ERROR] Pagination stopped:",
                error
            )

            break

    # Remove duplicate URLs
    unique_items = {}

    for item in all_items:
        unique_items[item["url"]] = item

    final_items = list(unique_items.values())

    print("\n============================================================")
    print(
        f" TOTAL UNIQUE PDFs FOUND: "
        f"{len(final_items)}"
    )
    print("============================================================")

    return final_items


# ============================================================
# LOAD OLD EXCEL RECORDS
# ============================================================

def load_old_records():

    old_records = {}

    if not EXCEL_FILE.exists():
        return old_records

    try:

        workbook = load_workbook(EXCEL_FILE)
        sheet = workbook.active

        headers = {}

        for index, cell in enumerate(
            sheet[1],
            start=1
        ):
            headers[cell.value] = index

        url_column = headers.get("PDF URL")
        path_column = headers.get("PDF Path")

        if not url_column:
            return old_records

        for row in sheet.iter_rows(
            min_row=2,
            values_only=True
        ):

            url = row[url_column - 1]

            if not url:
                continue

            path = ""

            if path_column:
                path = row[path_column - 1] or ""

            old_records[url] = {
                "path": str(path)
            }

    except Exception as error:

        print(
            "[WARNING] Could not read old Excel:",
            error
        )

    return old_records


# ============================================================
# DOWNLOAD PDF
# ============================================================

def download_pdf(page, pdf_url, output_path):

    try:

        time.sleep(REQUEST_DELAY_SECONDS + random.uniform(0, 0.35))
        response = page.context.request.get(
            pdf_url,
            timeout=90000
        )

        check_response_status(response.status, pdf_url)

        if not response.ok:

            return False, (
                f"HTTP {response.status}"
            )

        content = response.body()

        if not content.startswith(b"%PDF"):

            return False, "Not a valid PDF"

        output_path.write_bytes(content)

        if (
            output_path.exists()
            and output_path.stat().st_size > 500
        ):

            return True, ""

        return False, "Empty PDF file"

    except Exception as error:

        return False, str(error)


# ============================================================
# SAVE EXCEL REPORT
# ============================================================

def save_excel(records):

    workbook = Workbook()
    sheet = workbook.active

    sheet.title = "EPFO Circular Report"

    headers = [
        "PDF Name",
        "PDF URL",
        "PDF Path",
        "Downloaded Status",
        "Download Date & Time"
    ]

    sheet.append(headers)

    for record in records:

        sheet.append([
            record["name"],
            record["url"],
            record["path"],
            record["status"],
            record["date_time"]
        ])

    sheet.column_dimensions["A"].width = 60
    sheet.column_dimensions["B"].width = 90
    sheet.column_dimensions["C"].width = 90
    sheet.column_dimensions["D"].width = 25
    sheet.column_dimensions["E"].width = 25

    sheet.freeze_panes = "A2"

    for cell in sheet[1]:
        cell.font = cell.font.copy(bold=True)

    try:

        workbook.save(EXCEL_FILE)

        print(
            f"\n[OK] Excel saved: "
            f"{EXCEL_FILE.resolve()}"
        )

    except PermissionError:

        backup_file = Path(
            "EPFO_Circular_Report_"
            + datetime.now().strftime(
                "%Y%m%d_%H%M%S"
            )
            + ".xlsx"
        )

        workbook.save(backup_file)

        print(
            f"\n[WARNING] Excel was open."
            f"\nBackup saved: {backup_file.resolve()}"
        )


# ============================================================
# DOWNLOAD ALL CIRCULAR PDFs
# ============================================================

def download_all_circular_pdfs(page):

    print("\n============================================================")
    print(" DOWNLOADING ALL CIRCULAR PDFs")
    print("============================================================")

    CIRCULARS_FOLDER.mkdir(
        parents=True,
        exist_ok=True
    )

    old_records = load_old_records()

    circulars = collect_all_circulars(page)

    if not circulars:

        print("[WARNING] No PDF files found")
        return

    report_records = []

    used_names = Counter()

    newly_downloaded = 0
    already_downloaded = 0
    failed_downloads = 0

    for index, item in enumerate(
        circulars,
        start=1
    ):

        title = item["title"]
        pdf_url = item["url"]

        date_time = datetime.now().strftime(
            "%Y-%m-%d %H:%M:%S"
        )

        # Skip already downloaded URL
        if pdf_url in old_records:

            old_path = old_records[pdf_url]["path"]

            if old_path and Path(old_path).exists():

                print(
                    f"[{index}/{len(circulars)}] "
                    f"[SKIP] Already Downloaded: "
                    f"{title}"
                )

                report_records.append({
                    "name": Path(old_path).name,
                    "url": pdf_url,
                    "path": old_path,
                    "status": "Already Downloaded",
                    "date_time": date_time
                })

                already_downloaded += 1
                continue

        output_path = unique_file_path(
            CIRCULARS_FOLDER,
            title,
            used_names
        )

        print(
            f"[{index}/{len(circulars)}] "
            f"Downloading: {title}"
        )

        success, error = download_pdf(
            page,
            pdf_url,
            output_path
        )

        if success:

            status = "Newly Downloaded"
            newly_downloaded += 1

            print(
                f"    [SAVED] {output_path.name}"
            )

        else:

            status = "Failed"
            failed_downloads += 1

            print(
                f"    [FAILED] {error}"
            )

        report_records.append({
            "name": output_path.name,
            "url": pdf_url,
            "path": str(output_path.resolve()),
            "status": status,
            "date_time": date_time
        })

    save_excel(report_records)

    print("\n============================================================")
    print(" CIRCULAR DOWNLOAD COMPLETED")
    print("============================================================")

    print(
        f"PDFs found          : {len(circulars)}"
    )

    print(
        f"Newly Downloaded    : {newly_downloaded}"
    )

    print(
        f"Already Downloaded  : {already_downloaded}"
    )

    print(
        f"Failed              : {failed_downloads}"
    )

    print(
        f"Folder              : "
        f"{CIRCULARS_FOLDER.resolve()}"
    )

    print(
        f"Excel               : "
        f"{EXCEL_FILE.resolve()}"
    )


# ============================================================
# MAIN
# ============================================================

def main():

    print("\n")
    print("============================================================")
    print("                 EPFO AUTOMATION")
    print("============================================================")

    # --------------------------------------------------------
    # STEP 1: SAVE EPF ACT PDF
    # --------------------------------------------------------

    save_epf_act_as_pdf()

    # --------------------------------------------------------
    # STEP 2: OPEN BROWSER
    # --------------------------------------------------------

    with sync_playwright() as p:

        browser = p.chromium.launch(
            headless=HEADLESS
        )

        page = browser.new_page(
            viewport={
                "width": 1366,
                "height": 900
            }
        )

        page.set_default_timeout(
            TIMEOUT
        )

        # ----------------------------------------------------
        # STEP 3: OPEN EPFO WEBSITE
        # ----------------------------------------------------

        print("\n[STEP] Opening EPFO website...")

        page.goto(
            START_URL,
            wait_until="domcontentloaded",
            timeout=TIMEOUT
        )

        time.sleep(3)

        print("[OK] EPFO website opened")

        # ----------------------------------------------------
        # STEP 4: LEGAL FRAMEWORK
        # ----------------------------------------------------

        if not open_legal_framework(page):

            browser.close()
            return

        # ----------------------------------------------------
        # STEP 5: OPEN EPF ACT
        # ----------------------------------------------------

        if not open_epf_act(page):

            browser.close()
            return

        # ----------------------------------------------------
        # STEP 6: OPEN CIRCULARS
        # ----------------------------------------------------

        if not open_circulars(page):

            browser.close()
            return

        # ----------------------------------------------------
        # STEP 7: DOWNLOAD ALL CIRCULAR PDFs
        # ----------------------------------------------------

        download_all_circular_pdfs(page)

        # ----------------------------------------------------
        # FINISH
        # ----------------------------------------------------

        print("\n============================================================")
        print("                 ALL WORK COMPLETED")
        print("============================================================")

        print("\nEPF Act PDF Folder:")
        print(EPF_PDF_FOLDER.resolve())

        print("\nCircular PDFs Folder:")
        print(CIRCULARS_FOLDER.resolve())

        print("\nExcel Report:")
        print(EXCEL_FILE.resolve())

        print("\n============================================================")

        browser.close()


if __name__ == "__main__":
    guarded_main(main)
