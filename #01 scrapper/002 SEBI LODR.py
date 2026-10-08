import asyncio
import json
import os
import random
import re
from datetime import datetime
from pathlib import Path
from urllib.parse import unquote, urljoin, urlparse, parse_qs

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from playwright.async_api import async_playwright


SCRIPT_DIR = Path(__file__).resolve().parent
ENV_FILE = next(parent / ".env" for parent in Path(__file__).resolve().parents if (parent / "docker-compose.yml").is_file())
if ENV_FILE.exists():
    for line in ENV_FILE.read_text(encoding="utf-8").splitlines():
        if "=" in line and not line.strip().startswith("#"):
            key, value = line.split("=", 1)
            os.environ.setdefault(key.strip(), value.strip().strip('"'))

BASE_DIR = Path(os.environ.get("DOWNLOAD_BASE_PATH", SCRIPT_DIR)) / Path(__file__).stem
RUNTIME_DIR = BASE_DIR / "logs"
RUN_STATUS = RUNTIME_DIR / "run_status.json"
RUNTIME_LOG = RUNTIME_DIR / "runtime.log"
RESILIENCE_ATTEMPTS = int(os.environ.get("SCRAPER_MAX_ATTEMPTS", "3"))
REQUEST_DELAY = float(os.environ.get("SCRAPER_REQUEST_DELAY_SECONDS", "1.0"))
BLOCK_COOLDOWN = float(os.environ.get("SCRAPER_BLOCK_COOLDOWN_SECONDS", "300"))


class SiteAccessBlocked(RuntimeError):
    pass


BLOCK_DETECTED = ""


def check_response_status(status, url=""):
    global BLOCK_DETECTED
    if status in {403, 429}:
        BLOCK_DETECTED = f"HTTP {status} from {url}"
        raise SiteAccessBlocked(BLOCK_DETECTED)


def write_run_status(state, message="", attempt=0):
    RUNTIME_DIR.mkdir(parents=True, exist_ok=True)
    payload = {"script": Path(__file__).name, "state": state, "attempt": attempt,
               "message": str(message), "timestamp": datetime.now().astimezone().isoformat(timespec="seconds")}
    temporary = RUN_STATUS.with_suffix(".tmp")
    temporary.write_text(json.dumps(payload, indent=2), encoding="utf-8")
    temporary.replace(RUN_STATUS)
    with RUNTIME_LOG.open("a", encoding="utf-8") as stream:
        stream.write(f"{payload['timestamp']} | {state} | attempt={attempt} | {message}\n")


async def guarded_main(entrypoint):
    global BLOCK_DETECTED
    for attempt in range(1, RESILIENCE_ATTEMPTS + 1):
        BLOCK_DETECTED = ""
        write_run_status("running", attempt=attempt)
        try:
            result = await entrypoint()
            if BLOCK_DETECTED:
                raise SiteAccessBlocked(BLOCK_DETECTED)
            write_run_status("completed", attempt=attempt)
            return result
        except KeyboardInterrupt:
            write_run_status("stopped", "Stopped by user", attempt)
            raise
        except SiteAccessBlocked as error:
            write_run_status("blocked", error, attempt)
            if attempt == RESILIENCE_ATTEMPTS:
                raise
            await asyncio.sleep(BLOCK_COOLDOWN + random.uniform(0, 5))
        except Exception as error:
            write_run_status("error", error, attempt)
            if attempt == RESILIENCE_ATTEMPTS:
                raise
            await asyncio.sleep(min(120, 5 * (2 ** (attempt - 1))) + random.uniform(0, 2))
AUDIT_FILE = BASE_DIR / "SEBI_Audit.xlsx"
MAX_ITEMS = int(os.environ.get("SEBI_MAX_ITEMS", "0"))  # 0 = full run
HEADLESS = os.environ.get("SEBI_HEADLESS", "false").lower() in {"1", "true", "yes"}
BASE_URL = "https://www.sebi.gov.in"

SECTIONS = [
    ("Act", "Acts", "https://www.sebi.gov.in/sebiweb/home/HomeAction.do?doListing=yes&sid=1&ssid=1&smid=0"),
    ("Rule", "Rules", "https://www.sebi.gov.in/sebiweb/home/HomeAction.do?doListing=yes&sid=1&ssid=2&smid=0"),
    ("Regulation", "Regulations", "https://www.sebi.gov.in/sebiweb/home/HomeAction.do?doListing=yes&sid=1&ssid=3&smid=0"),
    ("General Order", "General_Orders", "https://www.sebi.gov.in/sebiweb/home/HomeAction.do?doListing=yes&sid=1&ssid=4&smid=0"),
    ("Guideline", "Guidelines", "https://www.sebi.gov.in/sebiweb/home/HomeAction.do?doListing=yes&sid=1&ssid=5&smid=0"),
    ("Master Circular", "Master_Circulars", "https://www.sebi.gov.in/sebiweb/home/HomeAction.do?doListing=yes&sid=1&ssid=6&smid=0"),
    ("Circular", "Circulars", "https://www.sebi.gov.in/sebiweb/home/HomeAction.do?doListing=yes&sid=1&ssid=7&smid=0"),
]


def safe_filename(value):
    value = re.sub(r"\s+", " ", value).strip()
    value = re.sub(r'[<>:"/\\|?*]', "_", value)
    return value[:170] or "SEBI_document"


def unique_path(folder, filename):
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / filename
    n = 1
    while path.exists():
        path = folder / f"{Path(filename).stem}_{n}{Path(filename).suffix}"
        n += 1
    return path


def write_audit(rows):
    BASE_DIR.mkdir(parents=True, exist_ok=True)
    if AUDIT_FILE.exists():
        workbook = load_workbook(AUDIT_FILE)
        if "Audit" in workbook.sheetnames:
            del workbook["Audit"]
    else:
        workbook = Workbook()
        if "Sheet" in workbook.sheetnames:
            del workbook["Sheet"]
    sheet = workbook.create_sheet("Audit")
    sheet.append([
        "Record No.", "Category", "Visible Date", "Visible Title", "Visible Row Metadata",
        "Document Label", "Source Page URL", "Detail Page URL", "Download URL",
        "PDF Path", "Status", "Error"
    ])
    for row in rows:
        sheet.append([
            row["record_no"], row["category"], row["date"], row["name"],
            row["visible_metadata"], row["document_label"], row["source_url"],
            row["detail_url"], row["download_url"], row["pdf_path"],
            row["status"], row["error"],
        ])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F5597")
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for col, width in {"A": 12, "B": 22, "C": 18, "D": 100, "E": 100, "F": 32,
                       "G": 80, "H": 80, "I": 100, "J": 115, "K": 18, "L": 50}.items():
        sheet.column_dimensions[col].width = width
    workbook.save(AUDIT_FILE)


async def find_pdfs(context, detail_url, folder, number, title):
    detail = await context.new_page()
    try:
        await detail.goto(detail_url, wait_until="domcontentloaded", timeout=45000)
        await detail.wait_for_timeout(1000)
        links = []
        for selector, attr in (("a[href]", "href"), ("iframe[src]", "src")):
            for element in await detail.query_selector_all(selector):
                value = await element.get_attribute(attr)
                if not value:
                    continue
                if "file=" in value:
                    value = unquote(parse_qs(urlparse(value).query).get("file", [""])[0])
                if ".pdf" not in value.lower() and "/sebi_data/attachdocs/" not in value.lower():
                    continue
                label = (await element.inner_text()).strip() if selector.startswith("a") else "PDF"
                links.append((label or "PDF", urljoin(BASE_URL, value)))
        links = list(dict.fromkeys(links))
        results = []
        for index, (label, pdf_url) in enumerate(links, 1):
            try:
                await asyncio.sleep(REQUEST_DELAY + random.uniform(0, 0.35))
                response = await context.request.get(pdf_url, timeout=60000)
                check_response_status(response.status, pdf_url)
                data = await response.body()
                if not response.ok or not data.startswith(b"%PDF"):
                    results.append({"document_label": label, "download_url": pdf_url,
                                    "pdf_path": "", "status": "Failed",
                                    "error": f"HTTP {response.status} / NOT PDF"})
                    continue
                suffix = "" if index == 1 else f"_{index:02d}"
                path = unique_path(folder, f"{number:04d}_{safe_filename(title)}{suffix}.pdf")
                path.write_bytes(data)
                results.append({"document_label": label, "download_url": pdf_url,
                                "pdf_path": str(path.resolve()), "status": "Downloaded", "error": ""})
            except Exception as exc:
                results.append({"document_label": label, "download_url": pdf_url,
                                "pdf_path": "", "status": "Failed", "error": str(exc)})
        return results or [{"document_label": "", "download_url": "", "pdf_path": "",
                            "status": "No PDF link", "error": "No PDF link found on detail page"}]
    except Exception as exc:
        return [{"document_label": "", "download_url": "", "pdf_path": "",
                 "status": "Failed", "error": f"{type(exc).__name__}: {exc}"}]
    finally:
        await detail.close()


async def process_section(context, page, category, folder_name, url, rows):
    folder = BASE_DIR / folder_name
    await page.goto(url, wait_until="domcontentloaded", timeout=45000)
    await page.wait_for_timeout(1800)
    seen = set()
    page_number = 1
    while True:
        table_rows = await page.query_selector_all("table tbody tr")
        page_items = []
        for row in table_rows:
            cells = await row.query_selector_all("td")
            if len(cells) < 2:
                continue
            date = (await cells[0].inner_text()).strip()
            link = await cells[1].query_selector("a")
            if not link:
                continue
            title = (await link.inner_text()).strip()
            href = await link.get_attribute("href")
            if href and title:
                cell_texts = [((await cell.inner_text()).strip()) for cell in cells]
                visible_metadata = " | ".join(cell_texts)
                page_items.append((title, date, visible_metadata, urljoin(BASE_URL, href)))
        for title, date, visible_metadata, detail_url in page_items:
            if detail_url in seen:
                continue
            seen.add(detail_url)
            detail_results = await find_pdfs(context, detail_url, folder, len(seen), title)
            for result in detail_results:
                rows.append({"record_no": len(seen), "name": title, "category": category, "date": date,
                             "visible_metadata": visible_metadata, "detail_url": detail_url,
                             **result, "source_url": url})
            write_audit(rows)
            print(f"{category} {len(seen)}: {title[:80]} -> {len(detail_results)} document(s)", flush=True)
            if MAX_ITEMS and len({x["record_no"] for x in rows if x["category"] == category}) >= MAX_ITEMS:
                return
        next_link = page.locator("a", has_text=re.compile(r"^Next", re.I)).last
        if not await next_link.count() or not await next_link.is_visible():
            return
        try:
            await next_link.click(timeout=15000)
            await page.wait_for_timeout(1600)
            page_number += 1
        except Exception:
            return


async def main():
    rows = []
    write_audit(rows)
    async with async_playwright() as p:
        browser = await p.chromium.launch(headless=HEADLESS)
        context = await browser.new_context()
        page = await context.new_page()
        for category, folder, url in SECTIONS:
            await process_section(context, page, category, folder, url, rows)
        await context.close()
        await browser.close()
    print(f"Regulations: {sum(x['category'] == 'Regulation' for x in rows)}")
    print(f"Master Circulars: {sum(x['category'] == 'Master Circular' for x in rows)}")
    print(f"Circulars: {sum(x['category'] == 'Circular' for x in rows)}")
    print(f"Audit: {AUDIT_FILE.resolve()}")


if __name__ == "__main__":
    asyncio.run(guarded_main(main))
